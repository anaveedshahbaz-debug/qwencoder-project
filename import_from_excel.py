"""
import_from_excel.py
=====================
Import all project data from your Excel monitoring file into the web app database.

Usage:
  1. Copy your Excel file (.xlsm) into this same folder
  2. Edit EXCEL_FILE below if the filename is different
  3. Run:  python import_from_excel.py
"""
import sys, os, sqlite3
from datetime import datetime

EXCEL_FILE = "2__CMD_-_Monitoring_of_Projects_-_August_2025.xlsm"
DB_FILE    = "nespak_pms.db"

os.chdir(os.path.dirname(os.path.abspath(__file__)))

try:
    from openpyxl import load_workbook
except ImportError:
    print("ERROR: openpyxl not installed.\nRun:  pip install openpyxl")
    sys.exit(1)

import app as appmod
appmod.init_db()

def detect_group(sheet_ref, project_name=""):
    if not sheet_ref: return "GENERAL"
    ref = str(sheet_ref).upper(); name = str(project_name).upper()
    if "-R-" in ref or "RRUDP" in name or "RUDA" in name or "RAVI" in name: return "RRUDP"
    if "-C-" in ref or "CBD" in name: return "CBD"
    if "-L-" in ref or "LDA" in name: return "LDA"
    if "-GPI-" in ref: return "GPI"
    if "-PMA-" in ref: return "PMA"
    if "-TMT-" in ref: return "TMT"
    if "-TEPA-" in ref: return "TEPA"
    if "-SA-" in ref: return "SA"
    return "GENERAL"

def safe_float(v):
    if v is None: return 0.0
    if isinstance(v,(int,float)): return float(v)
    try: return float(str(v).replace(",",""))
    except: return 0.0

def safe_date(v):
    if v is None: return None
    if isinstance(v, datetime): return v.strftime("%Y-%m-%d")
    if isinstance(v, str):
        for fmt in ("%Y-%m-%d","%d/%m/%Y","%m/%d/%Y"):
            try: return datetime.strptime(v, fmt).strftime("%Y-%m-%d")
            except: pass
    return None

def safe_str(v):
    return "" if v is None else str(v).strip()

PM_MAP = {
    "ZAHID MEHMOOD KHAN":"Zahid Mehmood Khan","ZAHID MAHMOOD KHAN":"Zahid Mehmood Khan",
    "IMRAN BADAR":"Imran Badar","WASEEM SAIF DAR":"Waseem Saif Dar",
    "MUHAMMAD ASLAM SOOMRO":"Muhammad Aslam Soomro","AHMAD NAVEED SHAHBAZ":"Ahmad Naveed Shahbaz",
    "ARSLAN ZAMIR":"Arslan Zamir","AFTAB AHMED":"Aftab Ahmed",
    "JAMSHAID FAISAL JANJUA":"Jamshaid Faisal Janjua","SANAULLAH HASHMI":"Sanaullah Hashmi",
    "SAJID HAMID":"Sajid Hamid",
}
def normalize_pm(name):
    if not name: return "Unknown"
    key = str(name).upper().strip()
    for k,v in PM_MAP.items():
        if k in key or key in k: return v
    return str(name).strip()

print("\n" + "="*60)
print("NESPAK PMS — Excel Data Import (Web Version)")
print("="*60)

if not os.path.exists(EXCEL_FILE):
    print(f"\nERROR: File not found: {EXCEL_FILE}")
    print("Copy your Excel file into this folder and try again.")
    sys.exit(1)

print(f"Loading: {EXCEL_FILE} ...")
wb = load_workbook(EXCEL_FILE, data_only=True, read_only=True)
ws = wb["Summary"]
print("Loaded.\n")

projects_data = []
for row in ws.iter_rows(values_only=True):
    if row[0] and isinstance(row[0],(int,float)) and row[1] and row[3]:
        projects_data.append(row)

print(f"Projects found: {len(projects_data)}")

inserted=skipped=errors=0
for row in projects_data:
    try:
        sheet_ref = safe_str(row[1]); job_no = safe_str(row[2])
        project_name = safe_str(row[3]); pm = normalize_pm(row[4])
        comm_date = safe_date(row[5]); comp_date = safe_date(row[6])
        amendment = safe_str(row[7]) or "None"
        if amendment in ("0",""): amendment="None"
        duration = safe_float(row[8]); nespak_fee = safe_float(row[9])
        salary_ce = safe_float(row[10]); direct_ce = safe_float(row[11])
        group = detect_group(sheet_ref, project_name)

        billing_upto = safe_float(row[16]) if len(row)>16 else 0.0
        billing_date = safe_date(row[17]) if len(row)>17 else None
        receipt_upto = safe_float(row[18]) if len(row)>18 else 0.0
        receipt_date = safe_date(row[19]) if len(row)>19 else None

        existing = appmod.query("SELECT id FROM projects WHERE sheet_ref=?",(sheet_ref,),one=True)
        if existing: skipped+=1; continue

        pid = appmod.execute("""INSERT INTO projects
            (sheet_ref,job_no,project_name,project_manager,project_group,
             comm_date,comp_date,amendment,duration_months,nespak_fee_mil,
             salary_cost_est,direct_cost_est,is_active)
            VALUES (?,?,?,?,?,?,?,?,?,?,?,?,1)""",
            (sheet_ref,job_no,project_name,pm,group,comm_date,comp_date,
             amendment,duration,nespak_fee,salary_ce,direct_ce))

        if billing_upto>0 or receipt_upto>0:
            snap_date = billing_date or receipt_date
            if snap_date:
                try:
                    dt=datetime.strptime(snap_date,"%Y-%m-%d"); snap_m,snap_y=dt.month,dt.year
                except: snap_m,snap_y=8,2025
            else: snap_m,snap_y=8,2025

            if billing_upto>0:
                appmod.execute("""INSERT INTO transactions
                    (project_id,txn_month,txn_year,txn_type,amount,
                     invoice_no,invoice_date,description)
                    VALUES (?,?,?,?,?,?,?,?)""",
                    (pid,snap_m,snap_y,"BILLING",billing_upto,
                     f"013/{job_no}/2025-26/01",f"{snap_y}-{snap_m:02d}-01",
                     "Imported - cumulative upto"))

            sal_actual=safe_float(row[22]) if len(row)>22 else 0.0
            dir_actual=safe_float(row[24]) if len(row)>24 else 0.0

            if receipt_upto>0:
                appmod.execute("""INSERT INTO transactions
                    (project_id,txn_month,txn_year,txn_type,amount,
                     cheque_amount,amount_received,description)
                    VALUES (?,?,?,?,?,?,?,?)""",
                    (pid,snap_m,snap_y,"RECEIPT",receipt_upto,
                     receipt_upto,receipt_upto,"Imported - cumulative upto"))

            if sal_actual>0 or dir_actual>0:
                appmod.execute("""INSERT INTO transactions
                    (project_id,txn_month,txn_year,txn_type,amount,
                     salary_actual,direct_actual,description)
                    VALUES (?,?,?,?,?,?,?,?)""",
                    (pid,snap_m,snap_y,"EXPENSE",sal_actual+dir_actual,
                     sal_actual,dir_actual,"Imported - actual expenditure upto"))

            appmod.rebuild_snapshot(pid, snap_m, snap_y)

        inserted+=1
        print(f"  OK  {sheet_ref:12s} {project_name[:50]}")
    except Exception as e:
        errors+=1
        print(f"  ERR row {row[0]}: {e}")

print(f"\n{'='*60}")
print(f"Import complete! Imported: {inserted}  Skipped: {skipped}  Errors: {errors}")
print("="*60)
print("\nNow run:  python app.py")
print("Then open: http://localhost:5000")
